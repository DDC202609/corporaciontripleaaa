#!/usr/bin/env python3
"""Importa operaciones históricas del libro Excel y genera su contabilidad.

Uso (primero siempre se recomienda simulación):
  PYTHONPATH="/ruta/a/site-packages" .venv/bin/python scripts/importar_historico_excel.py --dry-run
  PYTHONPATH="/ruta/a/site-packages" .venv/bin/python scripts/importar_historico_excel.py --apply

El importador es idempotente: VIN, factura, movimientos y claves de origen se
registran para que una segunda ejecución no duplique la información.
"""
import argparse
import os
import shutil
import sqlite3
import sys
import tempfile
from collections import defaultdict
from datetime import date, datetime, timedelta

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
import server  # noqa: E402

DEFAULT_BOOK = "/Users/erickmadrid/Library/Mobile Documents/com~apple~CloudDocs/Sistema de Gestión Corporación Triple AAA GPT.xlsm"


def text(value):
    return "" if value is None else str(value).strip()


def amount(value):
    if value in (None, "", "N/A", "#N/A"):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    cleaned = text(value).replace("L", "").replace(",", "").strip()
    try:
        return float(cleaned)
    except ValueError:
        return 0.0


def iso_date(value):
    if not value:
        return datetime.now().strftime("%Y-%m-%d")
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    raw = text(value)
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(raw[:10], pattern).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return datetime.now().strftime("%Y-%m-%d")


def normalized(value):
    return text(value).lower().replace(" ", "_").replace(".", "").replace("_", "")


def sheet_rows(workbook, name, header_row):
    sheet = workbook[name]
    headers = [normalized(cell.value) for cell in next(sheet.iter_rows(min_row=header_row, max_row=header_row))]
    for row_number, cells in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
        row = {headers[index]: cells[index] if index < len(cells) else None for index in range(len(headers)) if headers[index]}
        if any(value not in (None, "") for value in row.values()):
            yield row_number, row


def state(value):
    raw = text(value).lower()
    if "vend" in raw:
        return "Vendido"
    if "taller" in raw:
        return "En Taller"
    if "trans" in raw or "import" in raw:
        return "En Tránsito"
    if "nacional" in raw:
        return "Nacionalizado"
    if "reserv" in raw:
        return "Reservado"
    return "DPV"


def type_purchase(value):
    return "Importación" if "import" in text(value).lower() else "Compra Local"


def payment_type(value):
    raw = text(value).lower()
    if "tarjet" in raw:
        return "Tarjeta"
    if "deposit" in raw:
        return "Depósito"
    if "transfer" in raw or "tranferencia" in raw:
        return "Transferencia"
    return "Efectivo"


def expense_account(category, classification):
    value = (text(category) + " " + text(classification)).lower()
    if "planilla" in value or "salario" in value or "sueldo" in value:
        return "6201" if "admin" in value or "geren" in value else "6101"
    if "comision" in value:
        return "6402"
    if "public" in value:
        return "6401"
    if "combust" in value:
        return "6301"
    if "flete" in value:
        return "6403"
    if "seguro" in value:
        return "6205"
    if "alquiler" in value or "renta" in value:
        return "6206"
    if "servicio" in value or "energia" in value or "luz" in value:
        return "6204"
    if "papeler" in value or "insumo" in value:
        return "6203"
    if "honor" in value:
        return "6202"
    if "otro" in value:
        return "7201"
    return "6401" if "venta" in value else "6202"


def ensure_schema(connection):
    connection.execute('''CREATE TABLE IF NOT EXISTS importaciones_historicas(
        id INTEGER PRIMARY KEY AUTOINCREMENT, archivo TEXT NOT NULL UNIQUE,
        aplicado_en TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP)''')
    connection.execute('''CREATE TABLE IF NOT EXISTS registros_importacion_historica(
        id INTEGER PRIMARY KEY AUTOINCREMENT, importacion_id INTEGER NOT NULL REFERENCES importaciones_historicas(id),
        tipo TEXT NOT NULL, clave_origen TEXT NOT NULL, destino_tabla TEXT, destino_id INTEGER,
        UNIQUE(importacion_id,tipo,clave_origen))''')
    connection.execute('''CREATE TABLE IF NOT EXISTS gastos_operativos(
        id INTEGER PRIMARY KEY AUTOINCREMENT, fecha TEXT NOT NULL, categoria TEXT NOT NULL,
        clasificacion TEXT, proveedor TEXT, documento TEXT, subtotal REAL NOT NULL DEFAULT 0,
        isv REAL NOT NULL DEFAULT 0, total REAL NOT NULL DEFAULT 0, forma_pago TEXT, banco TEXT,
        origen_clave TEXT UNIQUE NOT NULL, creado_en TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP)''')


def tracked(connection, import_id, kind, key):
    return connection.execute('SELECT 1 FROM registros_importacion_historica WHERE importacion_id=? AND tipo=? AND clave_origen=?', (import_id, kind, key)).fetchone() is not None


def mark(connection, import_id, kind, key, table, row_id):
    connection.execute('''INSERT OR IGNORE INTO registros_importacion_historica(importacion_id,tipo,clave_origen,destino_tabla,destino_id)
        VALUES(?,?,?,?,?)''', (import_id, kind, key, table, row_id))


def provider(connection, name, credit=False):
    name = text(name) or "Proveedor histórico sin identificar"
    row = connection.execute('SELECT id FROM proveedores WHERE nombre=?', (name,)).fetchone()
    if row:
        return row['id']
    cursor = connection.execute('''INSERT INTO proveedores(nombre,condicion_pago,dias_credito,activo)
        VALUES(?,?,?,1)''', (name, "Crédito" if credit else "Contado", 30 if credit else 0))
    return cursor.lastrowid


def account_id(connection, code):
    row = connection.execute('SELECT id FROM cuentas_contables WHERE codigo=? AND activo=1', (code,)).fetchone()
    if not row:
        raise ValueError(f"No existe la cuenta contable requerida {code}")
    return row['id']


def import_data(book_path, database_path, apply=False):
    server.DB = database_path
    server.init_db()
    workbook = load_workbook(book_path, read_only=True, data_only=True, keep_vba=True)
    connection = server.db()
    ensure_schema(connection)
    source_name = os.path.abspath(book_path)
    source = connection.execute('SELECT id FROM importaciones_historicas WHERE archivo=?', (source_name,)).fetchone()
    if source:
        import_id = source['id']
    else:
        import_id = connection.execute('INSERT INTO importaciones_historicas(archivo) VALUES(?)', (source_name,)).lastrowid
    summary = defaultdict(int)
    flow_by_vin = {}
    for _, row in sheet_rows(workbook, 'Tabla_Contro_Flujo', 4):
        vin = text(row.get('vin')).upper()
        if vin and vin != '#N/A':
            flow_by_vin[vin] = row

    vehicles = {}
    for row_number, row in sheet_rows(workbook, 'Vehiculo', 5):
        vin = text(row.get('vin')).upper()
        if not vin or vin in ('#N/A', 'VIN'):
            continue
        key = f"Vehiculo:{row_number}:{vin}"
        existing = connection.execute('SELECT id FROM vehiculos WHERE vin=?', (vin,)).fetchone()
        if existing:
            vehicles[vin] = existing['id']; summary['vehiculos_omitidos'] += 1; mark(connection, import_id, 'vehiculo', key, 'vehiculos', existing['id']); continue
        plate = text(row.get('placa')).upper()
        if plate and connection.execute('SELECT 1 FROM vehiculos WHERE UPPER(TRIM(placa))=?', (plate,)).fetchone():
            plate = ''
            summary['placas_duplicadas_omitidas'] += 1
        cursor = connection.execute('''INSERT INTO vehiculos(vin,lote,tipo_vehiculo,marca,modelo,version,color,anio,placa,estado,ubicacion,
            precio_compra,precio_venta,fecha_adquisicion,proveedor,tipo_compra,observaciones)
            VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''', (
            vin, text(row.get('lote')), text(row.get('tipovehiculo')), text(row.get('marca')), text(row.get('modelo')),
            text(row.get('version')), text(row.get('color')), int(amount(row.get('año')) or 0) or None, plate,
            state(row.get('estatusgeneral')), 'AUTOLOTE', amount(row.get('costodecompra')), amount(row.get('precioventa')),
            iso_date(row.get('fechadecompra')), text(row.get('proveedor')), type_purchase(row.get('tipocompra')),
            'Importado desde historial Excel sin asiento contable original.'))
        vehicles[vin] = cursor.lastrowid; mark(connection, import_id, 'vehiculo', key, 'vehiculos', cursor.lastrowid); summary['vehiculos'] += 1

    # Incluye también unidades presentes en Kardex que no estaban en la ficha maestra.
    for row_number, row in sheet_rows(workbook, 'KARDEX', 5):
        vin = text(row.get('vin')).upper()
        if not vin or vin in vehicles:
            continue
        existing = connection.execute('SELECT id FROM vehiculos WHERE vin=?', (vin,)).fetchone()
        if existing:
            vehicles[vin] = existing['id']; continue
        plate = text(row.get('placa')).upper()
        if plate and connection.execute('SELECT 1 FROM vehiculos WHERE UPPER(TRIM(placa))=?', (plate,)).fetchone(): plate = ''
        cursor = connection.execute('''INSERT INTO vehiculos(vin,tipo_vehiculo,marca,modelo,version,color,anio,placa,estado,ubicacion,precio_compra,precio_venta,
            fecha_adquisicion,proveedor,tipo_compra,observaciones) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''', (
            vin, '', text(row.get('marca')), text(row.get('modelo')), text(row.get('version')), text(row.get('color')),
            int(amount(row.get('año')) or 0) or None, plate, state(row.get('estatus')), text(row.get('predio')) or 'AUTOLOTE',
            amount(row.get('costototal')), amount(row.get('preciodeventa')), iso_date(row.get('fecha')), text(row.get('proveedor')),
            type_purchase(row.get('tipodecompra')), 'Importado desde Kardex histórico.'))
        vehicles[vin] = cursor.lastrowid; summary['vehiculos_kardex'] += 1

    # Adquisiciones: las que tienen flujo conservan anticipo y CxP. Las demás
    # se contrapartidan con resultados acumulados para no inventar pagos.
    for vin, vehicle_id in vehicles.items():
        if connection.execute('SELECT 1 FROM adquisiciones WHERE vehiculo_id=?', (vehicle_id,)).fetchone():
            continue
        vehicle = connection.execute('SELECT * FROM vehiculos WHERE id=?', (vehicle_id,)).fetchone()
        flow = flow_by_vin.get(vin)
        total = amount(flow.get('costototal')) if flow else amount(vehicle['precio_compra'])
        if total <= 0: continue
        paid = min(amount(flow.get('montoapagado')), total) if flow else total
        pending = max(total - paid, 0)
        payment = payment_type(flow.get('formadepago')) if flow else 'Efectivo'
        bank = text(flow.get('banco')) if flow else ''
        date_value = iso_date(flow.get('fechadepago')) if flow else vehicle['fecha_adquisicion']
        supplier_id = provider(connection, vehicle['proveedor'], pending > 0)
        acquisition_id = connection.execute('''INSERT INTO adquisiciones(vehiculo_id,proveedor_id,fecha,costo_compra,anticipo,metodo_pago,condicion_pago,dias_credito,saldo,observaciones,necesita_reparacion,tipo_compra,banco)
            VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)''', (vehicle_id, supplier_id, date_value, total, paid, payment,
            'Crédito' if pending else 'Contado', 30 if pending else 0, pending,
            'Importación histórica de adquisición.', 0, vehicle['tipo_compra'] or 'Compra Local', bank)).lastrowid
        if flow:
            server.contabilizar_adquisicion(connection, acquisition_id)
        else:
            inventory = server.cuenta_inventario_adquisicion(connection.execute('SELECT * FROM adquisiciones WHERE id=?', (acquisition_id,)).fetchone())
            server.registrar_asiento(connection, date_value, f'Inventario histórico sin comprobante de pago · {vin}', 'adquisicion', acquisition_id,
                [{'cuenta_id': server.cuenta_contable_id(connection, inventory), 'debe': total}, {'cuenta_id': account_id(connection, '3103'), 'haber': total}], vehicle_id)
            # Evita que la sincronización posterior transforme una contrapartida
            # de migración en una salida de caja inexistente.
            server.registrar_movimiento_caja(connection, date_value, 'Contrapartida de migración', 'Migración', None, 0, 0,
                f'Adquisición histórica sin comprobante de pago · {vin}', 'adquisicion_anticipo', acquisition_id)
        summary['adquisiciones'] += 1

    for row_number, row in sheet_rows(workbook, 'Costo_Base', 5):
        vin = text(row.get('vin')).upper(); vehicle_id = vehicles.get(vin)
        if not vehicle_id: continue
        key = f"Costo_Base:{row_number}:{vin}"
        if tracked(connection, import_id, 'costeo', key): continue
        values = (amount(row.get('ajusteavalorcif')), amount(row.get('grúa')), amount(row.get('flete')), amount(row.get('clisvpagado')), amount(row.get('clstd')), amount(row.get('almacenaje')), amount(row.get('gastosaduanero')))
        connection.execute('''INSERT INTO costos_adquisicion(vehiculo_id,costo_exw,grua,flete,costo_estimado,ajuste_cif,isv_pagado,cl_std,almacenaje,gastos_aduaneros,estatus,ajuste_compra,placa_nacionalizacion)
            VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?) ON CONFLICT(vehiculo_id) DO UPDATE SET costo_exw=excluded.costo_exw,grua=excluded.grua,flete=excluded.flete,
            isv_pagado=excluded.isv_pagado,cl_std=excluded.cl_std,almacenaje=excluded.almacenaje,gastos_aduaneros=excluded.gastos_aduaneros,
            estatus=excluded.estatus,ajuste_compra=excluded.ajuste_compra,placa_nacionalizacion=excluded.placa_nacionalizacion''',
            (vehicle_id, amount(row.get('costoexw')), values[1], values[2], amount(row.get('costestimado')), values[0], values[3], values[4], values[5], values[6], text(row.get('estatus')) or 'Nacionalizado', values[0], text(row.get('placa'))))
        plate = text(row.get('placa')).upper()
        if plate and server.placa_disponible(connection, plate, vehicle_id): connection.execute('UPDATE vehiculos SET placa=? WHERE id=?', (plate, vehicle_id))
        mark(connection, import_id, 'costeo', key, 'costos_adquisicion', vehicle_id); summary['costeos'] += 1

    for row_number, row in sheet_rows(workbook, 'Gastos_Taller', 5):
        vin = text(row.get('vin')).upper(); vehicle_id = vehicles.get(vin)
        total = amount(row.get('total'))
        if not vehicle_id or total <= 0: continue
        key = f"Gastos_Taller:{row_number}:{vin}"
        if tracked(connection, import_id, 'gasto_taller', key): continue
        cursor = connection.execute('''INSERT INTO costos_vehiculo(vehiculo_id,fecha,concepto,categoria,monto,proveedor,documento,observaciones)
            VALUES(?,?,?,?,?,?,?,?)''', (vehicle_id, iso_date(row.get('fechaingresotaller')), text(row.get('tipodegasto')) or 'Taller',
            'Taller', total, text(row.get('proveedor')), text(row.get('factura')), f"{text(row.get('tallerasignado'))}: {text(row.get('descripción'))}"))
        mark(connection, import_id, 'gasto_taller', key, 'costos_vehiculo', cursor.lastrowid); summary['gastos_taller'] += 1

    payments_by_invoice = defaultdict(list)
    for row_number, row in sheet_rows(workbook, 'Tabla_de_Pago', 3):
        invoice = text(row.get('nodefactura'))
        if invoice: payments_by_invoice[invoice].append((row_number, row))
    for row_number, row in sheet_rows(workbook, 'Tabla_Venta', 4):
        invoice = text(row.get('notransacción'))
        vin = text(row.get('vin')).upper(); vehicle_id = vehicles.get(vin)
        if not invoice or not vehicle_id: continue
        key = f"Venta:{invoice}"
        if tracked(connection, import_id, 'venta', key) or connection.execute('SELECT 1 FROM ventas WHERE factura=?', (invoice,)).fetchone():
            summary['ventas_omitidas'] += 1; continue
        client_name = text(row.get('nombre')) or 'Consumidor Final'
        identity = text(row.get('identidad'))
        client = connection.execute('SELECT id FROM clientes WHERE nombre=? AND COALESCE(identidad,\'\')=? ORDER BY id LIMIT 1', (client_name, identity)).fetchone()
        if client: client_id = client['id']
        else: client_id = connection.execute('INSERT INTO clientes(nombre,identidad,rtn,telefono,direccion) VALUES(?,?,?,?,?)', (client_name, identity, text(row.get('rtn')), text(row.get('telefono')), text(row.get('dirección')))).lastrowid
        final = amount(row.get('preciofinal')) or amount(row.get('valortotal'))
        if final <= 0: continue
        number = server.siguiente_correlativo_plataforma(connection, 'VENTA')
        sale_id = connection.execute('''INSERT INTO ventas(vehiculo_id,cliente_id,fecha,precio,descuento,prima,saldo,estado,factura,correlativo,numero_transaccion,tipo_venta,forma_pago,financiera_banco,monto_financiado,transferencia,observaciones)
            VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''', (vehicle_id, client_id, iso_date(row.get('fecha')), final, amount(row.get('descuento')), amount(row.get('prima')), 0,
            'Facturada', invoice, number, f'V-{number:05d}', text(row.get('tipodeventa')) or 'Contado', text(row.get('formadepago')), text(row.get('financierabanco')),
            amount(row.get('montofinanciado')), 0, text(row.get('observaciones')) or f'Factura histórica {invoice}')).lastrowid
        paid = 0.0
        for payment_row, payment in payments_by_invoice.get(invoice, []):
            value = amount(payment.get('monto'))
            value = min(value, max(final - paid, 0))
            if value <= 0: continue
            paid += value
            reference = text(payment.get('nodereferencia'))
            if not reference or reference.upper() in ('N/A', 'NO APLICA') or connection.execute('SELECT 1 FROM pagos_venta WHERE referencia=?', (reference,)).fetchone(): reference = f'HIST-{invoice}-{payment_row}'
            connection.execute('''INSERT INTO pagos_venta(venta_id,tipo_pago,referencia,banco,monto,fecha) VALUES(?,?,?,?,?,?)''',
                (sale_id, payment_type(payment.get('formadepago')), reference, text(payment.get('banco')), value, iso_date(payment.get('fecha'))))
        financed = amount(row.get('montofinanciado'))
        financed = min(financed, max(final - paid, 0))
        if financed > 0:
            finance_name = text(row.get('financierabanco')) or 'Financiera histórica sin identificar'
            finance_id = provider(connection, finance_name, True)
            connection.execute('''INSERT INTO pagos_venta(venta_id,tipo_pago,referencia,financiera_id,financiera_nombre,monto,fecha) VALUES(?,?,?,?,?,?,?)''',
                (sale_id, 'Financiado', f'FIN-{invoice}', finance_id, finance_name, financed, iso_date(row.get('fecha'))))
            paid += financed
        if paid < final - 0.01:
            # La fuente no detalla el medio restante; se registra como efectivo
            # histórico para que el asiento de la venta quede completo.
            connection.execute('''INSERT INTO pagos_venta(venta_id,tipo_pago,referencia,monto,fecha) VALUES(?,?,?,?,?)''',
                (sale_id, 'Efectivo', f'EFECTIVO-HIST-{invoice}', round(final - paid, 2), iso_date(row.get('fecha'))))
        mark(connection, import_id, 'venta', key, 'ventas', sale_id); summary['ventas'] += 1

    for row_number, row in sheet_rows(workbook, 'Tabla_gastos_Pagos', 4):
        total = amount(row.get('total'))
        if total <= 0: continue
        key = f"Gasto:{row_number}"
        if tracked(connection, import_id, 'gasto_operativo', key): continue
        subtotal, tax = amount(row.get('subtotal')), amount(row.get('isv'))
        # El total pagado es el documento fuente para caja. Algunas filas del
        # libro tienen redondeos o subtotal desactualizado; se ajusta el
        # subtotal para que el asiento conserve el total real del documento.
        if tax > total:
            tax = 0
        subtotal = round(total - tax, 2)
        cursor = connection.execute('''INSERT INTO gastos_operativos(fecha,categoria,clasificacion,proveedor,documento,subtotal,isv,total,forma_pago,banco,origen_clave)
            VALUES(?,?,?,?,?,?,?,?,?,?,?)''', (iso_date(row.get('fechadepago')), text(row.get('tipopago')) or 'Gasto histórico', text(row.get('clasificación')),
            text(row.get('proveedor')), text(row.get('nodedocumento')), subtotal, tax, total, payment_type(row.get('formadepago')), text(row.get('banco')), key))
        expense_id = cursor.lastrowid; expense_code = expense_account(row.get('tipopago'), row.get('clasificación'))
        lines = [{'cuenta_id': account_id(connection, expense_code), 'debe': subtotal}]
        if tax: lines.append({'cuenta_id': account_id(connection, '1106'), 'debe': tax})
        lines.append({'cuenta_id': account_id(connection, '1101'), 'haber': total})
        description = f"Gasto histórico · {text(row.get('tipopago')) or 'Sin categoría'}"
        server.registrar_asiento(connection, iso_date(row.get('fechadepago')), description, 'gasto_operativo', expense_id, lines)
        medio, bank = server.medio_caja(payment_type(row.get('formadepago')), text(row.get('banco')))
        server.registrar_movimiento_caja(connection, iso_date(row.get('fechadepago')), 'Gasto operativo histórico', medio, bank, 0, total, description, 'gasto_operativo', expense_id)
        mark(connection, import_id, 'gasto_operativo', key, 'gastos_operativos', expense_id); summary['gastos_operativos'] += 1

    # Las ventas se contabilizan después de cargar todos sus pagos. Así se crean
    # CxC a financieras y movimientos de caja desde las operaciones reales.
    server.sincronizar_contabilidad_historica(connection)
    if apply:
        connection.commit()
    else:
        connection.rollback()
    connection.close()
    return summary


def main():
    parser = argparse.ArgumentParser(description='Migración histórica de Corporación Triple AAA')
    parser.add_argument('--archivo', default=DEFAULT_BOOK, help='Ruta del libro .xlsx/.xlsm')
    parser.add_argument('--base', default=os.path.join(ROOT, 'data', 'autolote.sqlite'), help='Ruta de la base SQLite')
    parser.add_argument('--apply', action='store_true', help='Confirma la aplicación de la migración')
    parser.add_argument('--dry-run', action='store_true', help='Simula sin guardar cambios')
    args = parser.parse_args()
    if args.apply == args.dry_run:
        parser.error('Use exactamente una opción: --dry-run o --apply')
    if not os.path.exists(args.archivo):
        parser.error(f'No existe el archivo: {args.archivo}')
    database_path = args.base
    temporary = None
    if args.dry_run:
        # Una simulación nunca abre la base operativa en modo escritura.
        temporary = tempfile.NamedTemporaryFile(prefix='autolote-import-', suffix='.sqlite', delete=False)
        temporary.close()
        shutil.copy2(args.base, temporary.name)
        database_path = temporary.name
    try:
        result = import_data(args.archivo, database_path, apply=args.apply)
    finally:
        if temporary and os.path.exists(temporary.name):
            os.unlink(temporary.name)
    print('SIMULACIÓN' if args.dry_run else 'MIGRACIÓN APLICADA')
    for key in sorted(result): print(f'{key}: {result[key]}')


if __name__ == '__main__':
    main()
