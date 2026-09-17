"""Valida las ventas y pagos históricos antes de importarlos."""
from collections import defaultdict
from pathlib import Path
import sqlite3
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DATABASE = ROOT / "data" / "autolote.sqlite"
SOURCE = Path("/Users/erickmadrid/Library/Mobile Documents/com~apple~CloudDocs/Autolote AAA/Sistema de Gestión Corporación Triple AAA V3 Version WEb.xlsm")


def invoice(value):
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def main():
    workbook = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True, keep_vba=True)
    sales_sheet = workbook["Tabla_Venta"]
    sales_headers = next(sales_sheet.iter_rows(min_row=4, max_row=4, values_only=True))
    sales_index = {str(value).strip(): i for i, value in enumerate(sales_headers) if value is not None}
    sales = []
    for row in sales_sheet.iter_rows(min_row=5, values_only=True):
        sale_invoice, vin = row[sales_index["No. Transacción"]], row[sales_index["Vin"]]
        if sale_invoice and vin:
            sales.append({
                "invoice": invoice(sale_invoice),
                "vin": str(vin).strip().upper(),
                "price": float(row[sales_index["Precio Final"]] or 0),
                "date": str(row[sales_index["Fecha"]])[:10],
                "name": str(row[sales_index["Nombre"]] or "").strip(),
                "type": str(row[sales_index["Tipo de Venta"]] or "").strip(),
                "method": str(row[sales_index["Medio de Pago"]] or "").strip(),
                "prima": float(row[sales_index["Prima"]] or 0),
                "financed": float(row[sales_index["Monto Financiado"]] or 0),
                "finance_name": str(row[sales_index["Financiera/Banco"]] or "").strip(),
            })
    payments_sheet = workbook["Tabla_de_Pago"]
    payment_headers = next(payments_sheet.iter_rows(min_row=3, max_row=3, values_only=True))
    payment_index = {str(value).strip(): i for i, value in enumerate(payment_headers) if value is not None}
    payments = defaultdict(list)
    for row in payments_sheet.iter_rows(min_row=4, values_only=True):
        payment_invoice, amount = row[payment_index["No. De Factura"]], row[payment_index["Monto"]]
        if payment_invoice not in (None, "") and amount not in (None, "") and float(amount or 0) > 0:
            payments[invoice(payment_invoice)].append({
                "type": str(row[payment_index["Forma de Pago"]] or "").strip(),
                "amount": float(amount),
                "date": str(row[payment_index["Fecha"]])[:10],
                "bank": str(row[payment_index["Banco"]] or "").strip(),
                "reference": str(row[payment_index["No. De Referencia"]] or "").strip(),
                "due": str(row[payment_index["Fecha de Vencimiento"]] or "")[:10],
            })
    connection = sqlite3.connect(DATABASE)
    known_vins = {row[0].upper() for row in connection.execute("SELECT vin FROM vehiculos")}
    existing_invoices = {str(row[0]) for row in connection.execute("SELECT factura FROM ventas WHERE factura IS NOT NULL")}
    connection.close()
    exact = [sale for sale in sales if sale["invoice"] in payments and abs(sum(item["amount"] for item in payments[sale["invoice"]]) - sale["price"]) <= .01]
    partial = [sale for sale in sales if sale["invoice"] in payments and abs(sum(item["amount"] for item in payments[sale["invoice"]]) - sale["price"]) > .01]
    no_payment = [sale for sale in sales if sale["invoice"] not in payments]
    print(f"Ventas fuente: {len(sales)}")
    duplicate_vins = defaultdict(list)
    for sale in sales:
        duplicate_vins[sale["vin"]].append(sale["invoice"])
    print(f"VIN repetidos en ventas: {[(vin, invoices) for vin, invoices in duplicate_vins.items() if len(invoices) > 1]}")
    print(f"Vehículos existentes: {sum(sale['vin'] in known_vins for sale in sales)}")
    print(f"VIN faltantes: {len([sale for sale in sales if sale['vin'] not in known_vins])}")
    print(f"Facturas ya registradas: {len([sale for sale in sales if sale['invoice'] in existing_invoices])}")
    print(f"Pagos que cuadran: {len(exact)}")
    print(f"Pagos descuadrados: {len(partial)}")
    print(f"Ventas sin detalle de pago: {len(no_payment)}")
    print("--- Pagos descuadrados ---")
    for sale in partial:
        total = sum(item["amount"] for item in payments[sale["invoice"]])
        print(sale["invoice"], sale["vin"], f"venta={sale['price']:.2f}", f"pagos={total:.2f}")
    print("--- Ventas sin detalle de pago ---")
    for sale in no_payment:
        print(sale["invoice"], sale["vin"], f"venta={sale['price']:.2f}", sale["type"], sale["method"])
    print("--- Pagos sin venta ---")
    sale_invoices = {sale["invoice"] for sale in sales}
    for payment_invoice, items in payments.items():
        if payment_invoice not in sale_invoices:
            print(payment_invoice, f"pagos={sum(item['amount'] for item in items):.2f}")
    print("--- Cobertura de financiamiento de ventas descuadradas/sin pago ---")
    for sale in partial + no_payment:
        total_payment = sum(item["amount"] for item in payments.get(sale["invoice"], []))
        balance = round(sale["price"] - total_payment, 2)
        print(
            sale["invoice"], sale["vin"], f"saldo={balance:.2f}",
            f"prima={sale['prima']:.2f}", f"financiado={sale['financed']:.2f}",
            sale["type"], sale["method"], sale["finance_name"], sale["name"],
        )


if __name__ == "__main__":
    main()
