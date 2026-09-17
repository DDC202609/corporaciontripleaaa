#!/usr/bin/env python3
"""Carga las ventas históricas y sus cobros desde el libro corporativo.

Reglas acordadas para esta carga:
* se omite la factura 100015 y se conserva la 100017 para su VIN duplicado;
* los números de factura históricos se conservan como referencia, mientras los
  correlativos V-xxxxx los genera la plataforma;
* los pagos sin detalle se llevan a efectivo;
* el financiamiento de junio, julio y agosto se cobra en BAC; septiembre queda
  pendiente en CxC. Los pagarés siempre permanecen en CxC del cliente.

El script se detiene si ya hay facturas históricas importadas y genera un CSV
de auditoría para validar cada venta y cobro cargado.
"""
import argparse
import csv
import shutil
import sqlite3
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]

SOURCE = Path("/Users/erickmadrid/Library/Mobile Documents/com~apple~CloudDocs/Autolote AAA/Sistema de Gestión Corporación Triple AAA V3 Version WEb.xlsm")
DATABASE = ROOT / "data" / "autolote.sqlite"
REPORT = ROOT / "data" / "reporte_migracion_ventas_20260915.csv"
OMIT_INVOICES = {"100015"}
CUENTAS = {
    "caja_bancos": "1101", "cxc": "1102", "inventario_dpv": "1103",
    "inventario_taller": "1104", "inventario_transito": "1105",
    "ingresos_venta": "4101", "costo_ventas": "5101",
}


def db():
    connection = sqlite3.connect(DATABASE)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys=ON")
    return connection


def account(c, key):
    row = c.execute("SELECT id FROM cuentas_contables WHERE codigo=? AND activo=1", (CUENTAS[key],)).fetchone()
    if not row:
        raise ValueError(f"No existe la cuenta contable {CUENTAS[key]}")
    return row["id"]


def inventory_account(state):
    return {"DPV": "inventario_dpv", "En Taller": "inventario_taller", "En Tránsito": "inventario_transito"}.get(state, "inventario_dpv")


def next_sale_number(c):
    row = c.execute("SELECT ultimo_numero FROM correlativos_plataforma WHERE tipo='VENTA'").fetchone()
    number = int(row["ultimo_numero"] or 0) + 1 if row else 1
    c.execute("INSERT INTO correlativos_plataforma(tipo,ultimo_numero) VALUES('VENTA',?) ON CONFLICT(tipo) DO UPDATE SET ultimo_numero=excluded.ultimo_numero", (number,))
    return number


def journal(c, fecha, descripcion, reference_type, reference_id, lines, vehicle_id):
    if c.execute("SELECT 1 FROM asientos_contables WHERE referencia_tipo=? AND referencia_id=?", (reference_type, reference_id)).fetchone():
        return
    debit = round(sum(float(line.get("debe") or 0) for line in lines), 2)
    credit = round(sum(float(line.get("haber") or 0) for line in lines), 2)
    if abs(debit - credit) > .01:
        raise ValueError(f"Asiento descuadrado: {descripcion}")
    seat_id = c.execute("INSERT INTO asientos_contables(fecha,descripcion,referencia_tipo,referencia_id) VALUES(?,?,?,?)", (fecha, descripcion, reference_type, reference_id)).lastrowid
    for line in lines:
        c.execute("""INSERT INTO partidas(fecha,descripcion,referencia_tipo,referencia_id,cuenta_id,vehiculo_id,debe,haber,asiento_id)
            VALUES(?,?,?,?,?,?,?,?,?)""", (fecha, descripcion, reference_type, reference_id, line["cuenta_id"], vehicle_id, round(float(line.get("debe") or 0), 2), round(float(line.get("haber") or 0), 2), seat_id))


def cash_movement(c, fecha, tipo, medio, banco, entrada, descripcion, reference_type, reference_id):
    c.execute("""INSERT INTO movimientos_caja(fecha,tipo,medio,banco,entrada,salida,descripcion,referencia_tipo,referencia_id)
        VALUES(?,?,?,?,?,0,?,?,?)""", (fecha, tipo, medio, banco, round(entrada, 2), descripcion, reference_type, reference_id))


def consolidated_cost(c, vehicle_id):
    vehicle = c.execute("SELECT precio_compra FROM vehiculos WHERE id=?", (vehicle_id,)).fetchone()
    base = float(vehicle["precio_compra"] or 0)
    costing = c.execute("SELECT ajuste_compra,grua,flete,isv_pagado,cl_std,almacenaje,gastos_aduaneros FROM costos_adquisicion WHERE vehiculo_id=?", (vehicle_id,)).fetchone()
    if costing:
        base += sum(float(costing[key] or 0) for key in ("ajuste_compra", "grua", "flete", "isv_pagado", "cl_std", "almacenaje", "gastos_aduaneros"))
    extras = c.execute("SELECT COALESCE(SUM(monto),0) valor FROM costos_vehiculo WHERE vehiculo_id=?", (vehicle_id,)).fetchone()
    return base + float(extras["valor"] or 0)


def account_sale(c, sale_id, vehicle):
    sale = c.execute("""SELECT ve.*,cl.nombre cliente_nombre FROM ventas ve
        LEFT JOIN clientes cl ON cl.id=ve.cliente_id WHERE ve.id=?""", (sale_id,)).fetchone()
    payments = c.execute("SELECT * FROM pagos_venta WHERE venta_id=? ORDER BY id", (sale_id,)).fetchall()
    lines = []
    for payment in payments:
        amount = float(payment["monto"] or 0)
        if payment["tipo_pago"] in ("Financiado", "Pagaré"):
            lines.append({"cuenta_id": account(c, "cxc"), "debe": amount})
            name = payment["financiera_nombre"] if payment["tipo_pago"] == "Financiado" else sale["cliente_nombre"]
            c.execute("""INSERT INTO cuentas_por_cobrar(venta_id,financiera_id,financiera_cliente_id,cliente_id,financiera_nombre,fecha,fecha_vencimiento,monto_original,saldo,estado)
                VALUES(?,?,?,?,?,?,?,?,?,?)""", (sale_id, payment["financiera_id"], payment["financiera_cliente_id"], sale["cliente_id"] if payment["tipo_pago"] == "Pagaré" else None, name or "Cliente no especificado", sale["fecha"], payment["fecha_vencimiento"], amount, amount, "Pendiente"))
        else:
            lines.append({"cuenta_id": account(c, "caja_bancos"), "debe": amount})
            medium = "Efectivo" if payment["tipo_pago"] == "Efectivo" else "Banco"
            bank = None if medium == "Efectivo" else (payment["banco"] or "Banco no especificado")
            cash_movement(c, payment["fecha"], "Cobro de venta", medium, bank, amount, f"Venta {sale['numero_transaccion']}", "pago_venta", payment["id"])
    lines.append({"cuenta_id": account(c, "ingresos_venta"), "haber": float(sale["precio"] or 0)})
    journal(c, sale["fecha"], f"Venta de vehículo · {sale['numero_transaccion']}", "venta_ingreso", sale_id, lines, vehicle["id"])
    cost = consolidated_cost(c, vehicle["id"])
    if cost:
        journal(c, sale["fecha"], f"Costo de venta · {sale['numero_transaccion']}", "venta_costo", sale_id, [
            {"cuenta_id": account(c, "costo_ventas"), "debe": cost},
            {"cuenta_id": account(c, inventory_account(vehicle["estado"])), "haber": cost},
        ], vehicle["id"])


def clean(value):
    return "" if value is None else str(value).strip()


def number(value):
    if value in (None, "", "N/A", "-"):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    return float(clean(value).replace("L", "").replace(",", "").replace(" ", "") or 0)


def invoice(value):
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return clean(value)


def iso(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    raw = clean(value)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(raw[:10], fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError(f"Fecha inválida: {raw!r}")


def payment_type(value):
    raw = clean(value).lower()
    if "tarjet" in raw:
        return "Tarjeta"
    if "deposit" in raw:
        return "Depósito"
    if "transfer" in raw or "tranferencia" in raw:
        return "Transferencia"
    return "Efectivo"


def normalized(value):
    return "".join(clean(value).lower().replace("á", "a").replace("é", "e").replace("í", "i").replace("ó", "o").replace("ú", "u").split())


def rows(book):
    sales_sheet = book["Tabla_Venta"]
    sales_headers = {clean(value): index for index, value in enumerate(next(sales_sheet.iter_rows(min_row=4, max_row=4, values_only=True))) if value is not None}
    sales = []
    for row in sales_sheet.iter_rows(min_row=5, values_only=True):
        sale_invoice = invoice(row[sales_headers["No. Transacción"]])
        vin = clean(row[sales_headers["Vin"]]).upper()
        if not sale_invoice or not vin:
            continue
        sales.append({
            "invoice": sale_invoice, "vin": vin,
            "fecha": iso(row[sales_headers["Fecha"]]),
            "nombre": clean(row[sales_headers["Nombre"]]) or "Consumidor final",
            "rtn": clean(row[sales_headers["RTN"]]),
            "identidad": clean(row[sales_headers["Identidad"]]),
            "telefono": clean(row[sales_headers["Telefono"]]),
            "direccion": clean(row[sales_headers["Dirección"]]),
            "precio": round(number(row[sales_headers["Precio Final"]]), 2),
            "descuento": round(number(row[sales_headers["Descuento"]]), 2),
            "tipo": clean(row[sales_headers["Tipo de Venta"]]),
            "prima": round(number(row[sales_headers["Prima"]]), 2),
            "financiado": round(number(row[sales_headers["Monto Financiado"]]), 2),
            "financiera": clean(row[sales_headers["Financiera/Banco"]]),
            "observaciones": clean(row[sales_headers["Observaciones"]]),
        })
    payment_sheet = book["Tabla_de_Pago"]
    payment_headers = {clean(value): index for index, value in enumerate(next(payment_sheet.iter_rows(min_row=3, max_row=3, values_only=True))) if value is not None}
    payments = defaultdict(list)
    for row in payment_sheet.iter_rows(min_row=4, values_only=True):
        source_invoice = invoice(row[payment_headers["No. De Factura"]])
        amount = number(row[payment_headers["Monto"]])
        if source_invoice and amount > 0:
            payments[source_invoice].append({
                "fecha": iso(row[payment_headers["Fecha"]]),
                "tipo": payment_type(row[payment_headers["Forma de Pago"]]),
                "banco": clean(row[payment_headers["Banco"]]),
                "referencia": clean(row[payment_headers["No. De Referencia"]]),
                "monto": round(amount, 2),
            })
    return sales, payments


def financial_client(c, name):
    target = normalized(name)
    for row in c.execute("SELECT id,nombre FROM clientes WHERE tipo='Financiera' AND activo=1"):
        if normalized(row["nombre"]) == target:
            return row
    raise ValueError(f"No existe la financiera activa {name!r}")


def customer(c, sale):
    identity = sale["identidad"]
    rtn = sale["rtn"]
    # Los ceros del Excel son valores genéricos y no identifican a una persona.
    identity_key = identity if identity and set(identity) != {"0"} else ""
    rtn_key = rtn if rtn and set(rtn) != {"0"} else ""
    if identity_key:
        row = c.execute("SELECT id FROM clientes WHERE identidad=? ORDER BY id LIMIT 1", (identity_key,)).fetchone()
        if row:
            return row["id"]
    if rtn_key:
        row = c.execute("SELECT id FROM clientes WHERE rtn=? ORDER BY id LIMIT 1", (rtn_key,)).fetchone()
        if row:
            return row["id"]
    row = c.execute("SELECT id FROM clientes WHERE lower(trim(nombre))=lower(trim(?)) AND tipo='Cliente' ORDER BY id LIMIT 1", (sale["nombre"],)).fetchone()
    if row:
        return row["id"]
    return c.execute('''INSERT INTO clientes(nombre,identidad,rtn,telefono,direccion,tipo,condicion_pago,activo)
        VALUES(?,?,?,?,?,'Cliente','Contado',1)''', (sale["nombre"], identity_key or None, rtn_key or None, sale["telefono"] or None, sale["direccion"] or None)).lastrowid


def safe_reference(c, original, invoice_number, sequence):
    candidate = clean(original)
    if not candidate or candidate.upper() in {"N/A", "NA", "0", "-"}:
        candidate = f"HIST-{invoice_number}-{sequence:02d}"
    if c.execute("SELECT 1 FROM pagos_venta WHERE referencia=?", (candidate,)).fetchone():
        candidate = f"HIST-{invoice_number}-{sequence:02d}"
    return candidate


def insert_payment(c, sale_id, payment, sequence):
    reference = None if payment["tipo"] == "Efectivo" else safe_reference(c, payment.get("referencia"), payment["invoice"], sequence)
    c.execute('''INSERT INTO pagos_venta(venta_id,tipo_pago,referencia,financiera_cliente_id,financiera_nombre,banco,monto,fecha,fecha_vencimiento)
        VALUES(?,?,?,?,?,?,?,?,?)''', (sale_id, payment["tipo"], reference, payment.get("financiera_cliente_id"), payment.get("financiera_nombre"), payment.get("banco") or None, payment["monto"], payment["fecha"], payment.get("fecha_vencimiento")))


def import_sales(apply=False):
    book = load_workbook(SOURCE, read_only=True, data_only=True, keep_vba=True)
    sales, source_payments = rows(book)
    sales = [sale for sale in sales if sale["invoice"] not in OMIT_INVOICES]
    if len(sales) != 77:
        raise ValueError(f"Se esperaban 77 ventas después de excluir 100015; se encontraron {len(sales)}")
    c = db()
    existing = c.execute("SELECT factura FROM ventas WHERE factura IS NOT NULL").fetchall()
    existing_invoices = {clean(row["factura"]) for row in existing}
    conflicts = [sale["invoice"] for sale in sales if sale["invoice"] in existing_invoices]
    if conflicts:
        raise ValueError("Ya existen facturas históricas; no se importó nada: " + ", ".join(conflicts))
    vins = {row["vin"]: row for row in c.execute("SELECT id,vin,estado,ubicacion FROM vehiculos")}
    missing = [sale["vin"] for sale in sales if sale["vin"] not in vins]
    if missing:
        raise ValueError("VIN sin vehículo maestro: " + ", ".join(missing))
    audit = []
    if not apply:
        print(f"Simulación correcta: {len(sales)} ventas. No se modificó la base.")
        return
    backup = DATABASE.with_name("autolote-antes-migracion-ventas-20260915.sqlite")
    if not backup.exists():
        shutil.copy2(DATABASE, backup)
    try:
        c.execute("BEGIN")
        for sale in sales:
            vehicle = vins[sale["vin"]]
            client_id = customer(c, sale)
            is_financed = normalized(sale["tipo"]) == "financiado" and sale["financiado"] > 0
            finance_amount = min(sale["financiado"], sale["precio"]) if is_financed else 0.0
            direct_amount = round(sale["precio"] - finance_amount, 2)
            number = next_sale_number(c)
            numero_venta = f"V-{number:05d}"
            cur = c.execute('''INSERT INTO ventas(vehiculo_id,cliente_id,fecha,precio,descuento,prima,saldo,estado,factura,correlativo,numero_transaccion,tipo_venta,forma_pago,financiera_banco,monto_financiado,transferencia,observaciones)
                VALUES(?,?,?,?,?,?,0,'Facturada',?,?,?,?,?,?,?,?,?)''', (
                vehicle["id"], client_id, sale["fecha"], sale["precio"], sale["descuento"], direct_amount,
                sale["invoice"], number, numero_venta, sale["tipo"] or "Contado", "Migración histórica",
                sale["financiera"] or None, finance_amount, 0,
                f"Migración histórica. Factura origen {sale['invoice']}. {sale['observaciones']}".strip()))
            sale_id = cur.lastrowid
            payment_sequence = 0
            used_source = []
            source_total = round(sum(item["monto"] for item in source_payments.get(sale["invoice"], [])), 2)
            # Los pagos detallados se usan hasta el componente no financiado. Si
            # exceden ese componente, se preserva el total comercial y se deja
            # constancia en el reporte para reclasificación/validación posterior.
            if direct_amount and 0 < source_total <= direct_amount + 0.01:
                used_source = source_payments[sale["invoice"]]
                for item in used_source:
                    payment_sequence += 1
                    insert_payment(c, sale_id, {**item, "invoice": sale["invoice"]}, payment_sequence)
                remainder = round(direct_amount - source_total, 2)
                if remainder > 0.01:
                    payment_sequence += 1
                    insert_payment(c, sale_id, {"invoice": sale["invoice"], "tipo": "Efectivo", "monto": remainder, "fecha": sale["fecha"]}, payment_sequence)
            elif direct_amount:
                payment_sequence += 1
                insert_payment(c, sale_id, {"invoice": sale["invoice"], "tipo": "Efectivo", "monto": direct_amount, "fecha": sale["fecha"]}, payment_sequence)
            finance_name = ""
            if finance_amount:
                payment_sequence += 1
                if normalized(sale["financiera"]) in {"pagare", "pagarecliente"}:
                    finance_name = sale["nombre"]
                    finance_payment = {"invoice": sale["invoice"], "tipo": "Pagaré", "monto": finance_amount, "fecha": sale["fecha"], "financiera_nombre": finance_name}
                else:
                    finance = financial_client(c, sale["financiera"])
                    finance_name = finance["nombre"]
                    finance_payment = {"invoice": sale["invoice"], "tipo": "Financiado", "monto": finance_amount, "fecha": sale["fecha"], "financiera_cliente_id": finance["id"], "financiera_nombre": finance_name}
                insert_payment(c, sale_id, finance_payment, payment_sequence)
            # Crea las pólizas de ingreso/costo y el auxiliar CxC antes de que
            # el vehículo pase a Vendido, para descargar DPV correctamente.
            account_sale(c, sale_id, vehicle)
            c.execute("UPDATE vehiculos SET estado='Vendido' WHERE id=?", (vehicle["id"],))
            c.execute('''INSERT INTO movimientos_vehiculo(vehiculo_id,fecha,tipo,estado_anterior,estado_nuevo,ubicacion_anterior,ubicacion_nueva,referencia,observaciones)
                VALUES(?,?,?,?,?,?,?,?,?)''', (vehicle["id"], sale["fecha"], "Venta histórica migrada", vehicle["estado"], "Vendido", vehicle["ubicacion"], vehicle["ubicacion"], sale["invoice"], f"Factura histórica {sale['invoice']} · {numero_venta}"))
            cxc_status = "No aplica"
            if finance_amount and normalized(sale["financiera"]) not in {"pagare", "pagarecliente"} and datetime.fromisoformat(sale["fecha"]).month in {6, 7, 8}:
                cxc = c.execute("SELECT id,saldo FROM cuentas_por_cobrar WHERE venta_id=? AND financiera_nombre=?", (sale_id, finance_name)).fetchone()
                if not cxc:
                    raise ValueError(f"No se generó CxC de la venta {sale['invoice']}")
                cobro_id = c.execute('''INSERT INTO cobros_cuentas_por_cobrar(cuenta_por_cobrar_id,fecha,tipo_pago,banco,referencia,monto)
                    VALUES(?,?,?,'BAC Credomatic',?,?)''', (cxc["id"], sale["fecha"], "Transferencia", f"MIG-COBRO-{sale['invoice']}", finance_amount)).lastrowid
                c.execute("UPDATE cuentas_por_cobrar SET saldo=0,estado='Cobrado' WHERE id=?", (cxc["id"],))
                cash_movement(c, sale["fecha"], "Cobro de financiera histórica", "Banco", "BAC Credomatic", finance_amount, f"Cobro financiado · {numero_venta}", "cobro_cxc", cobro_id)
                journal(c, sale["fecha"], f"Cobro financiado histórico · {numero_venta}", "cobro_cxc", cobro_id, [
                    {"cuenta_id": account(c, "caja_bancos"), "debe": finance_amount},
                    {"cuenta_id": account(c, "cxc"), "haber": finance_amount},
                ], vehicle["id"])
                cxc_status = "Cobrado en BAC"
            elif finance_amount:
                cxc_status = "Pendiente CxC"
            issue = ""
            if source_total > direct_amount + 0.01:
                issue = f"Pagos fuente ({source_total:.2f}) exceden pago directo ({direct_amount:.2f}); se cargó efectivo por el total directo."
            elif source_total and source_total < direct_amount - 0.01:
                issue = f"Pagos fuente parciales ({source_total:.2f}); complemento cargado en efectivo ({direct_amount-source_total:.2f})."
            audit.append({"factura_origen": sale["invoice"], "venta": numero_venta, "vin": sale["vin"], "precio": sale["precio"], "pago_directo": direct_amount, "financiado": finance_amount, "financiera": finance_name, "cxc": cxc_status, "pagos_fuente": source_total, "observacion": issue})
        c.commit()
    except Exception:
        c.rollback()
        raise
    finally:
        c.close()
    with REPORT.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(audit[0]))
        writer.writeheader()
        writer.writerows(audit)
    print(f"Importación completada: {len(audit)} ventas. Auditoría: {REPORT}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--apply", action="store_true", help="aplica la migración; sin esta opción solo simula")
    args = parser.parse_args()
    import_sales(apply=args.apply)
