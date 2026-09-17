#!/usr/bin/env python3
"""Completa las adquisiciones históricas pendientes desde Vehiculos.xlsx.

Usa el Costo Acumulado como valor de inventario y acredita CxP proveedores.
Las adquisiciones que ya fueron migradas se conservan sin duplicarse.
"""
import os
import sys
import re
from datetime import datetime

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
import server

ARCHIVO = "/Users/erickmadrid/Library/Mobile Documents/com~apple~CloudDocs/Autolote AAA/Vehiculos.xlsx"
HOJA = "Vehiculo (2)"
CUENTA_CXP = "2101"


def texto(valor):
    return "" if valor is None else str(valor).strip()


def fecha_iso(valor):
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    fecha = re.sub(r"/+", "/", texto(valor))
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(fecha, formato).date().isoformat()
        except ValueError:
            continue
    raise ValueError(f"Fecha inválida: {fecha}")


def proveedor_id(conexion, nombre):
    row = conexion.execute("SELECT id FROM proveedores WHERE nombre=?", (nombre,)).fetchone()
    if row:
        return row["id"]
    cur = conexion.execute(
        """INSERT INTO proveedores(nombre,condicion_pago,dias_credito,tipo,activo)
        VALUES(?,'Contado',0,'Proveedor',1)""", (nombre,)
    )
    return cur.lastrowid


def cuenta_id(conexion, codigo):
    row = conexion.execute(
        "SELECT id FROM cuentas_contables WHERE codigo=? AND activo=1", (codigo,)
    ).fetchone()
    if not row:
        raise ValueError(f"No existe la cuenta activa {codigo}.")
    return row["id"]


def main():
    libro = load_workbook(ARCHIVO, read_only=True, data_only=True)
    hoja = libro[HOJA]
    encabezados = {texto(c.value): i for i, c in enumerate(hoja[5])}
    requeridos = ["VIN", "Fecha_de_Compra", "Proveedor", "Tipo_Compra", "Costo Acumulado"]
    faltantes = [campo for campo in requeridos if campo not in encabezados]
    if faltantes:
        raise ValueError(f"Faltan columnas: {', '.join(faltantes)}")

    tipos = {"Compra Local": "Compra local", "Importación": "Importación", "Cambio": "Cambio"}
    registros = []
    for fila in hoja.iter_rows(min_row=6, values_only=True):
        vin = texto(fila[encabezados["VIN"]]).upper()
        if not vin:
            continue
        tipo_original = texto(fila[encabezados["Tipo_Compra"]])
        registro = {
            "vin": vin,
            "fecha": fecha_iso(fila[encabezados["Fecha_de_Compra"]]),
            "proveedor": texto(fila[encabezados["Proveedor"]]),
            "tipo": tipos.get(tipo_original),
            "costo": float(fila[encabezados["Costo Acumulado"]] or 0),
        }
        if not registro["proveedor"] or not registro["tipo"] or registro["costo"] <= 0:
            raise ValueError(f"Datos incompletos para el VIN {vin}.")
        registros.append(registro)
    if len(registros) != 188:
        raise ValueError(f"Se esperaban 188 vehículos y se encontraron {len(registros)}.")

    server.init_db()
    conexion = server.db()
    creadas = 0
    omitidas = 0
    total = 0.0
    try:
        cxp = cuenta_id(conexion, CUENTA_CXP)
        inv_dpv = cuenta_id(conexion, "1103")
        inv_transito = cuenta_id(conexion, "1105")
        for registro in registros:
            vehiculo = conexion.execute(
                "SELECT * FROM vehiculos WHERE vin=?", (registro["vin"],)
            ).fetchone()
            if not vehiculo:
                raise ValueError(f"No existe el vehículo maestro {registro['vin']}.")
            if conexion.execute("SELECT 1 FROM adquisiciones WHERE vehiculo_id=?", (vehiculo["id"],)).fetchone():
                omitidas += 1
                continue
            proveedor = proveedor_id(conexion, registro["proveedor"])
            nota = "Migración histórica. Contrapartida contable: CxP proveedores."
            cur = conexion.execute(
                """INSERT INTO adquisiciones
                (vehiculo_id,proveedor_id,fecha,costo_compra,anticipo,metodo_pago,
                 condicion_pago,dias_credito,saldo,observaciones,necesita_reparacion,
                 tipo_compra,placa_cambio,banco)
                VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (vehiculo["id"], proveedor, registro["fecha"], registro["costo"], 0,
                 "Efectivo", "Contado", 0, 0, nota, 0, registro["tipo"], None, None),
            )
            conexion.execute(
                """UPDATE vehiculos
                SET precio_compra=?,proveedor=?,fecha_adquisicion=?,tipo_compra=?
                WHERE id=?""",
                (registro["costo"], registro["proveedor"], registro["fecha"],
                 registro["tipo"], vehiculo["id"]),
            )
            estado = server.recalcular_estado(
                conexion, vehiculo["id"], "Etapa calculada al registrar adquisición histórica"
            )
            server.add_movimiento(
                conexion, vehiculo["id"], registro["fecha"], "Registro de adquisición histórica",
                vehiculo["estado"], estado, vehiculo["ubicacion"], vehiculo["ubicacion"],
                referencia="CxP proveedores",
                observaciones=f"Compra {registro['tipo']}: {registro['proveedor']}. {nota}",
            )
            inventario = inv_transito if registro["tipo"] == "Importación" else inv_dpv
            server.registrar_asiento(
                conexion, registro["fecha"],
                f"Adquisición histórica · {registro['vin']} · CxP proveedores",
                "migracion_adquisicion_socios", cur.lastrowid,
                [{"cuenta_id": inventario, "debe": registro["costo"]},
                 {"cuenta_id": cxp, "haber": registro["costo"]}],
                vehiculo["id"],
            )
            creadas += 1
            total += registro["costo"]
        conexion.commit()
    except Exception:
        conexion.rollback()
        raise
    finally:
        conexion.close()
        libro.close()
    print(f"Adquisiciones nuevas: {creadas}")
    print(f"Adquisiciones ya migradas: {omitidas}")
    print(f"Total nuevo contra CxP proveedores: {total:,.2f}")


if __name__ == "__main__":
    main()
