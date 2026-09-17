"""Carga inicial de gastos confidenciales desde Confi.xlsx.

El importe contable se calcula siempre como subtotal + ISV. Cuando el campo
"Total" de la fuente no coincide, se conserva la diferencia en observaciones
para que la carga sea auditable sin desbalancear el asiento.
"""
from __future__ import annotations

import sqlite3
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DB = ROOT / "data" / "autolote.sqlite"
SOURCE = Path("/Users/erickmadrid/Library/Mobile Documents/com~apple~CloudDocs/Autolote AAA/Confi.xlsx")
SHEET = "Tabla_gastos_Confi"
SOURCE_KEY = "migracion_gastos_confidenciales_confi"
REFERENCE_TYPE = "gasto_operativo"


def text(value):
    return str(value or "").strip()


def number(value):
    if value in (None, ""):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    cleaned = text(value).replace("L", "").replace(",", "").replace(" ", "")
    if cleaned in {"", "-"}:
        return 0.0
    return float(cleaned)


def iso_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return datetime.strptime(text(value), "%d/%m/%y").date().isoformat()


def source_rows():
    if not SOURCE.exists():
        raise FileNotFoundError(f"No se encontró el archivo fuente: {SOURCE}")
    book = load_workbook(SOURCE, read_only=True, data_only=True)
    try:
        sheet = book[SHEET]
        rows = []
        # La fila 4 es el encabezado y la información inicia en la 5.
        for excel_row, values in enumerate(sheet.iter_rows(min_row=5, values_only=True), 5):
            concept = text(values[2])
            if not concept:
                continue
            subtotal = round(number(values[5]), 2)
            isv = round(number(values[6]), 2)
            total_source = round(number(values[7]), 2)
            total = round(subtotal + isv, 2)
            if total <= 0:
                raise ValueError(f"Fila {excel_row}: el total calculado debe ser mayor que cero.")
            rows.append({
                "excel_row": excel_row,
                "concept": concept,
                "document": text(values[3]),
                "provider": text(values[4]),
                "subtotal": subtotal,
                "isv": isv,
                "total_source": total_source,
                "total": total,
                "date": iso_date(values[9]),
                "classification": text(values[11]),
            })
        return rows
    finally:
        book.close()


def account_id(cursor, code):
    row = cursor.execute(
        "SELECT id FROM cuentas_contables WHERE codigo=? AND activo=1 AND acepta_movimiento=1",
        (code,),
    ).fetchone()
    if not row:
        raise ValueError(f"No existe la cuenta activa {code}.")
    return row[0]


def migrate(apply=False):
    rows = source_rows()
    connection = sqlite3.connect(DB)
    connection.row_factory = sqlite3.Row
    cursor = connection.cursor()
    try:
        concepts = {row["nombre"]: row for row in cursor.execute("""
            SELECT cg.id, cg.nombre, cg.cuenta_id
            FROM conceptos_gasto cg
            JOIN cuentas_contables cc ON cc.id = cg.cuenta_id
            WHERE cg.activo=1 AND cc.activo=1 AND cc.acepta_movimiento=1
        """)}
        missing = sorted({row["concept"] for row in rows} - set(concepts))
        if missing:
            raise ValueError("Conceptos sin configuración contable: " + ", ".join(missing))
        existing = {row[0] for row in cursor.execute(
            "SELECT origen_clave FROM gastos_operativos WHERE origen_clave LIKE ?",
            (SOURCE_KEY + ":%",),
        )}
        pending = [row for row in rows if f"{SOURCE_KEY}:{row['excel_row']}" not in existing]
        total = round(sum(row["total"] for row in pending), 2)
        mismatches = [row for row in pending if abs(row["total"] - row["total_source"]) > 0.01]
        print(f"Registros fuente: {len(rows)}")
        print(f"Pendientes: {len(pending)}")
        print(f"Total confidencial calculado: {total:,.2f}")
        print(f"Filas con diferencia entre total fuente y calculado: {len(mismatches)}")
        for row in mismatches:
            print(f"  Fila {row['excel_row']}: {row['provider']} · fuente {row['total_source']:,.2f} · calculado {row['total']:,.2f}")
        if not apply:
            print("Revisión finalizada. Use --apply para registrar la carga.")
            return

        bank_account = account_id(cursor, "1101")
        isv_account = account_id(cursor, "1106")
        cursor.execute("BEGIN")
        for row in pending:
            concept = concepts[row["concept"]]
            provider = cursor.execute("SELECT id FROM proveedores WHERE nombre=?", (row["provider"],)).fetchone()
            note = f"Migración de gasto confidencial desde Confi.xlsx, fila {row['excel_row']}."
            if abs(row["total"] - row["total_source"]) > 0.01:
                note += f" Total fuente L {row['total_source']:,.2f}; total calculado L {row['total']:,.2f}."
            expense = cursor.execute("""
                INSERT INTO gastos_operativos(
                    concepto_id, proveedor_id, fecha, categoria, clasificacion, proveedor,
                    documento, subtotal, isv, total, forma_pago, banco, origen_clave,
                    referencia, observaciones
                ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
            """, (
                concept["id"], provider["id"] if provider else None, row["date"], row["concept"],
                f"Confidencial · {row['classification']}", row["provider"], row["document"] or None,
                row["subtotal"], row["isv"], row["total"], "Transferencia", "BAC Credomatic",
                f"{SOURCE_KEY}:{row['excel_row']}", None, note,
            ))
            expense_id = expense.lastrowid
            description = f"Pago de gasto confidencial · {row['concept']} · {row['provider']}"
            journal = cursor.execute("""
                INSERT INTO asientos_contables(fecha, descripcion, referencia_tipo, referencia_id)
                VALUES(?,?,?,?)
            """, (row["date"], description, REFERENCE_TYPE, expense_id))
            journal_id = journal.lastrowid
            cursor.execute("""
                INSERT INTO partidas(fecha,descripcion,referencia_tipo,referencia_id,cuenta_id,vehiculo_id,debe,haber,asiento_id)
                VALUES(?,?,?,?,?,?,?,?,?)
            """, (row["date"], description, REFERENCE_TYPE, expense_id, concept["cuenta_id"], None,
                  row["subtotal"], 0, journal_id))
            if row["isv"]:
                cursor.execute("""
                    INSERT INTO partidas(fecha,descripcion,referencia_tipo,referencia_id,cuenta_id,vehiculo_id,debe,haber,asiento_id)
                    VALUES(?,?,?,?,?,?,?,?,?)
                """, (row["date"], description, REFERENCE_TYPE, expense_id, isv_account, None,
                      row["isv"], 0, journal_id))
            cursor.execute("""
                INSERT INTO partidas(fecha,descripcion,referencia_tipo,referencia_id,cuenta_id,vehiculo_id,debe,haber,asiento_id)
                VALUES(?,?,?,?,?,?,?,?,?)
            """, (row["date"], description, REFERENCE_TYPE, expense_id, bank_account, None,
                  0, row["total"], journal_id))
            cursor.execute("""
                INSERT INTO movimientos_caja(fecha,tipo,medio,banco,entrada,salida,descripcion,referencia_tipo,referencia_id)
                VALUES(?,?,?,?,?,?,?,?,?)
            """, (row["date"], "Pago de gasto confidencial", "Banco", "BAC Credomatic", 0,
                  row["total"], description, REFERENCE_TYPE, expense_id))
        connection.commit()
        print(f"Carga aplicada: {len(pending)} registros por L {total:,.2f}.")
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()


if __name__ == "__main__":
    migrate("--apply" in sys.argv)
