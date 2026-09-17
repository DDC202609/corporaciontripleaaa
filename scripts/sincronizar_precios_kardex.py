"""Sincroniza Precio de Venta desde KARDEX, usando VIN como llave única."""
from pathlib import Path
import sqlite3
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DATABASE = ROOT / "data" / "autolote.sqlite"
SOURCE = Path("/Users/erickmadrid/Library/Mobile Documents/com~apple~CloudDocs/Autolote AAA/Sistema de Gestión Corporación Triple AAA V3 Version WEb.xlsm")


def main():
    workbook = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True, keep_vba=True)
    sheet = workbook["KARDEX"]
    headers = next(sheet.iter_rows(min_row=5, max_row=5, values_only=True))
    index = {str(value).strip(): position for position, value in enumerate(headers) if value is not None}
    prices = {
        str(row[index["VIN"]]).strip().upper(): round(float(row[index["Precio_de_Venta"]]), 2)
        for row in sheet.iter_rows(min_row=6, values_only=True)
        if row[index["VIN"]]
        and row[index["Precio_de_Venta"]] not in (None, "")
        and float(row[index["Precio_de_Venta"]] or 0) > 0
    }
    connection = sqlite3.connect(DATABASE)
    connection.row_factory = sqlite3.Row
    vehicles = connection.execute("""
        SELECT id, vin, estado, precio_venta, ubicacion
        FROM vehiculos WHERE COALESCE(precio_venta, 0)>0
        ORDER BY vin
    """).fetchall()
    missing = [row["vin"] for row in vehicles if row["vin"].upper() not in prices]
    if missing:
        raise RuntimeError("No se actualizó nada. VIN sin Precio_de_Venta en KARDEX: " + ", ".join(missing))
    corrections = [
        (row, prices[row["vin"].upper()])
        for row in vehicles
        if abs(float(row["precio_venta"] or 0) - prices[row["vin"].upper()]) > 0.01
    ]
    try:
        for row, price in corrections:
            connection.execute("UPDATE vehiculos SET precio_venta=? WHERE id=?", (price, row["id"]))
            connection.execute("""
                INSERT INTO movimientos_vehiculo
                (vehiculo_id, fecha, tipo, estado_anterior, estado_nuevo,
                 ubicacion_anterior, ubicacion_nueva, referencia, observaciones)
                VALUES (?, '2026-09-15', 'Precio de venta corregido', ?, ?, ?, ?,
                        'KARDEX', ?)
            """, (
                row["id"], row["estado"], row["estado"], row["ubicacion"], row["ubicacion"],
                f"Precio de venta sincronizado desde KARDEX: {price:.2f}",
            ))
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()
    print(f"Precios corregidos desde KARDEX: {len(corrections)}")
    for row, price in corrections:
        print(f"{row['vin']}: {row['precio_venta']:.2f} -> {price:.2f}")


if __name__ == "__main__":
    main()
