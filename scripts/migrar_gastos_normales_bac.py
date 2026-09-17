"""Migra Tabla_gastos_Pagos como pagos de gastos desde BAC Credomatic.

Uso:
  python scripts/migrar_gastos_normales_bac.py          # revisión, no escribe
  python scripts/migrar_gastos_normales_bac.py --apply  # registra la migración

Los importes se calculan siempre como subtotal + ISV. Esto corrige el único
registro histórico cuyo total no coincidía con sus componentes.
"""
from __future__ import annotations

import sqlite3
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DB = ROOT / "data" / "autolote.sqlite"
SOURCE = Path(
    "/Users/erickmadrid/Library/Mobile Documents/com~apple~CloudDocs/Autolote AAA/"
    "Sistema de Gestión Corporación Triple AAA V3 Version WEb.xlsm"
)
SHEET = "Tabla_gastos_Pagos"
REFERENCE_TYPE = "gasto_operativo"
SOURCE_KEY = "migracion_gastos_normales"


def number(value) -> float:
    return round(float(value or 0), 2)


def text(value) -> str:
    return str(value or "").strip()


def date_value(value) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    return datetime.fromisoformat(str(value)).date().isoformat()


def read_rows():
    book = load_workbook(SOURCE, read_only=True, data_only=True, keep_vba=True)
    sheet = book[SHEET]
    rows = []
    for excel_row, row in enumerate(sheet.iter_rows(min_row=5, values_only=True), start=5):
        concept = text(row[2])
        if not concept:
            continue
        subtotal = number(row[5])
        isv = number(row[6])
        total = round(subtotal + isv, 2)
        if total <= 0:
            raise ValueError(f"Fila {excel_row}: el total calculado debe ser mayor que cero.")
        original_total = number(row[7])
        note = (
            f"Migración histórica {SHEET}, fila {excel_row}. "
            f"Pago/origen Excel: {text(row[8]) or 'No indicado'} · {text(row[10]) or 'No indicado'}."
        )
        if abs(original_total - total) > 0.01:
            note += f" Total de Excel {original_total:.2f}; se usó subtotal + ISV = {total:.2f}."
        rows.append({
            "excel_row": excel_row,
            "vin": text(row[1]),
            "concept": concept,
            "document": text(row[3]),
            "provider": text(row[4]) or "Proveedor no especificado",
            "subtotal": subtotal,
            "isv": isv,
            "total": total,
            "date": date_value(row[9]),
            "classification": text(row[11]),
            "note": note,
        })
    return rows


def account_id(connection, code: str) -> int:
    row = connection.execute(
        "SELECT id FROM cuentas_contables WHERE codigo=? AND activo=1 AND acepta_movimiento=1", (code,)
    ).fetchone()
    if not row:
        raise ValueError(f"No existe una cuenta activa que acepte movimientos: {code}.")
    return row[0]


def insert_journal(connection, expense_id: int, item: dict, expense_account_id: int, bac_account_id: int, isv_account_id: int):
    description = f"Pago de gasto · {item['concept']} · {item['provider']}"
    cursor = connection.execute(
        "INSERT INTO asientos_contables(fecha,descripcion,referencia_tipo,referencia_id) VALUES(?,?,?,?)",
        (item["date"], description, REFERENCE_TYPE, expense_id),
    )
    journal_id = cursor.lastrowid
    lines = [(expense_account_id, item["subtotal"], 0)]
    if item["isv"]:
        lines.append((isv_account_id, item["isv"], 0))
    lines.append((bac_account_id, 0, item["total"]))
    if round(sum(line[1] for line in lines), 2) != round(sum(line[2] for line in lines), 2):
        raise ValueError(f"Fila {item['excel_row']}: el asiento no cuadra.")
    for account, debit, credit in lines:
        connection.execute(
            """INSERT INTO partidas(fecha,descripcion,referencia_tipo,referencia_id,cuenta_id,vehiculo_id,debe,haber,asiento_id)
               VALUES(?,?,?,?,?,?,?,?,?)""",
            (item["date"], description, REFERENCE_TYPE, expense_id, account, None, debit, credit, journal_id),
        )
    connection.execute(
        """INSERT INTO movimientos_caja(fecha,tipo,medio,banco,entrada,salida,descripcion,referencia_tipo,referencia_id)
           VALUES(?,?,?,?,?,?,?,?,?)""",
        (item["date"], "Pago de gasto", "Banco", "BAC Credomatic", 0, item["total"], description, REFERENCE_TYPE, expense_id),
    )


def migrate(apply: bool):
    rows = read_rows()
    connection = sqlite3.connect(DB)
    connection.row_factory = sqlite3.Row
    concepts = {
        row["nombre"]: row
        for row in connection.execute(
            """SELECT cg.id,cg.nombre,cg.cuenta_id FROM conceptos_gasto cg
               JOIN cuentas_contables cc ON cc.id=cg.cuenta_id
               WHERE cg.activo=1 AND cc.activo=1 AND cc.acepta_movimiento=1"""
        )
    }
    missing = sorted({row["concept"] for row in rows} - set(concepts))
    if missing:
        raise ValueError("Conceptos sin cuenta contable activa: " + ", ".join(missing))
    existing = {
        row[0]
        for row in connection.execute("SELECT origen_clave FROM gastos_operativos WHERE origen_clave LIKE ?", (SOURCE_KEY + ":%",))
    }
    pending = [row for row in rows if f"{SOURCE_KEY}:{row['excel_row']}" not in existing]
    print(f"Registros fuente: {len(rows)}")
    print(f"Ya migrados: {len(rows) - len(pending)}")
    print(f"Pendientes: {len(pending)}")
    print(f"Total pendiente: {sum(row['total'] for row in pending):,.2f}")
    if not apply:
        connection.close()
        print("Revisión finalizada. Use --apply para registrar la migración.")
        return
    bac_account_id = account_id(connection, "1101")
    isv_account_id = account_id(connection, "1106")
    try:
        connection.execute("BEGIN")
        for item in pending:
            provider = connection.execute("SELECT id FROM proveedores WHERE nombre=?", (item["provider"],)).fetchone()
            cursor = connection.execute(
                """INSERT INTO gastos_operativos(concepto_id,proveedor_id,fecha,categoria,clasificacion,proveedor,
                   documento,subtotal,isv,total,forma_pago,banco,origen_clave,referencia,observaciones)
                   VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    concepts[item["concept"]]["id"], provider["id"] if provider else None,
                    item["date"], item["concept"], item["classification"] or "Gasto operativo", item["provider"],
                    item["document"] or None, item["subtotal"], item["isv"], item["total"], "Transferencia",
                    "BAC Credomatic", f"{SOURCE_KEY}:{item['excel_row']}", None, item["note"],
                ),
            )
            insert_journal(
                connection, cursor.lastrowid, item, concepts[item["concept"]]["cuenta_id"], bac_account_id, isv_account_id
            )
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()
    print(f"Migración aplicada: {len(pending)} gastos por {sum(row['total'] for row in pending):,.2f}.")


if __name__ == "__main__":
    migrate("--apply" in sys.argv)
