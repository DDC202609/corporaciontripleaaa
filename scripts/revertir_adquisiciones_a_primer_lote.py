#!/usr/bin/env python3
"""Revierte adquisiciones históricas fuera del primer lote de 19 vehículos."""
import os
import sys

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
import server

ARCHIVO = "/Users/erickmadrid/Library/Mobile Documents/com~apple~CloudDocs/Autolote AAA/Lote de Costos 18.xlsx"


def main():
    hoja = load_workbook(ARCHIVO, read_only=True, data_only=True)["Hoja1"]
    vinos_primer_lote = {
        str(fila[1] or "").strip().upper()
        for fila in hoja.iter_rows(min_row=4, values_only=True)
        if str(fila[1] or "").strip()
    }
    if len(vinos_primer_lote) != 19:
        raise ValueError("El archivo de referencia no contiene exactamente 19 VIN.")

    conexion = server.db()
    try:
        rows = conexion.execute("""SELECT a.id, a.vehiculo_id, v.vin
            FROM adquisiciones a JOIN vehiculos v ON v.id=a.vehiculo_id
            WHERE COALESCE(a.observaciones,'') LIKE 'Migración histórica.%'""").fetchall()
        revertir = [row for row in rows if row["vin"].upper() not in vinos_primer_lote]
        adquisiciones = [row["id"] for row in revertir]
        vehiculos = [row["vehiculo_id"] for row in revertir]
        if not adquisiciones:
            print("No hay adquisiciones fuera del primer lote para revertir.")
            return

        marks_a = ",".join("?" for _ in adquisiciones)
        marks_v = ",".join("?" for _ in vehiculos)
        asientos = [row["id"] for row in conexion.execute(
            f"""SELECT id FROM asientos_contables
                WHERE referencia_tipo='migracion_adquisicion_socios'
                AND referencia_id IN ({marks_a})""", adquisiciones
        ).fetchall()]

        conexion.execute("BEGIN")
        if asientos:
            marks_asientos = ",".join("?" for _ in asientos)
            conexion.execute(f"DELETE FROM partidas WHERE asiento_id IN ({marks_asientos})", asientos)
            conexion.execute(f"DELETE FROM asientos_contables WHERE id IN ({marks_asientos})", asientos)
        conexion.execute(f"""DELETE FROM movimientos_vehiculo
            WHERE vehiculo_id IN ({marks_v})
            AND tipo IN ('Registro de adquisición histórica','Cambio automático de etapa')""", vehiculos)
        conexion.execute(f"DELETE FROM adquisiciones WHERE id IN ({marks_a})", adquisiciones)
        conexion.execute(f"""UPDATE vehiculos
            SET precio_compra=0, proveedor=NULL, fecha_adquisicion=NULL,
                tipo_compra=NULL, estado='En Tránsito'
            WHERE id IN ({marks_v})""", vehiculos)
        conexion.commit()
        print(f"Adquisiciones revertidas: {len(adquisiciones)}")
        print(f"Asientos revertidos: {len(asientos)}")
    except Exception:
        conexion.rollback()
        raise
    finally:
        conexion.close()


if __name__ == "__main__":
    main()
