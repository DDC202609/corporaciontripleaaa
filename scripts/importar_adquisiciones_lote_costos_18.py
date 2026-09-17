#!/usr/bin/env python3
"""Registra las 19 adquisiciones históricas del lote de costos.

Cada adquisición usa la contrapartida temporal Cuentas por Pagar - Socios.
No crea pagos, cuentas por pagar a proveedores ni reclasificaciones.
"""
import os
import sys
from datetime import datetime

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
import server

ARCHIVO = "/Users/erickmadrid/Library/Mobile Documents/com~apple~CloudDocs/Autolote AAA/Lote de Costos 18.xlsx"
HOJA = "Hoja1"
LIMITE = 19
CUENTA_CXP_SOCIOS = "2104"


def texto(valor):
    return "" if valor is None else str(valor).strip()


def fecha_iso(valor):
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    return datetime.fromisoformat(texto(valor)).date().isoformat()


def proveedor_id(conexion, nombre):
    proveedor = conexion.execute(
        "SELECT id FROM proveedores WHERE nombre=?", (nombre,)
    ).fetchone()
    if proveedor:
        return proveedor["id"]
    # Se conserva el proveedor del archivo como maestro, con la condición inicial acordada.
    cur = conexion.execute(
        """INSERT INTO proveedores
        (nombre,condicion_pago,dias_credito,tipo,activo)
        VALUES(?,'Contado',0,'Proveedor',1)""",
        (nombre,),
    )
    return cur.lastrowid


def cuenta_id(conexion, codigo):
    cuenta = conexion.execute(
        "SELECT id FROM cuentas_contables WHERE codigo=? AND activo=1", (codigo,)
    ).fetchone()
    if not cuenta:
        raise ValueError(f"No existe la cuenta contable {codigo}.")
    return cuenta["id"]


def asegurar_cxp_socios(conexion):
    conexion.execute(
        """INSERT OR IGNORE INTO cuentas_contables
        (codigo,cuenta,tipo,grupo,naturaleza,acepta_movimiento,requiere_activo,
         requiere_centro_costo,activo)
        VALUES(?,?,?,?,?,?,?,?,?)""",
        (CUENTA_CXP_SOCIOS, "Cuentas por Pagar - Socios", "Pasivo",
         "Pasivo Corriente", "Acreedora", 1, 0, 1, 1),
    )


def main():
    libro = load_workbook(ARCHIVO, read_only=True, data_only=True)
    hoja = libro[HOJA]
    encabezados = {texto(c.value): i for i, c in enumerate(hoja[3])}
    requeridos = ["VIN", "Fecha_de_Compra", "Proveedor", "Tipo_Compra", "Costo Acumulado"]
    faltantes = [nombre for nombre in requeridos if nombre not in encabezados]
    if faltantes:
        raise ValueError(f"Faltan columnas requeridas: {', '.join(faltantes)}")

    registros = []
    for fila in hoja.iter_rows(min_row=4, values_only=True):
        vin = texto(fila[encabezados["VIN"]]).upper()
        if not vin:
            continue
        registros.append({
            "vin": vin,
            "fecha": fecha_iso(fila[encabezados["Fecha_de_Compra"]]),
            "proveedor": texto(fila[encabezados["Proveedor"]]),
            "tipo": texto(fila[encabezados["Tipo_Compra"]]),
            "costo": float(fila[encabezados["Costo Acumulado"]] or 0),
        })
    registros = registros[:LIMITE]
    if len(registros) != LIMITE:
        raise ValueError(f"Se esperaban {LIMITE} registros y se encontraron {len(registros)}.")

    tipos = {"Compra Local": "Compra local", "Importación": "Importación", "Cambio": "Cambio"}
    for registro in registros:
        if not registro["proveedor"] or registro["costo"] <= 0:
            raise ValueError(f"Datos incompletos para el VIN {registro['vin']}.")
        if registro["tipo"] not in tipos:
            raise ValueError(f"Tipo de compra inválido para {registro['vin']}: {registro['tipo']}")
        registro["tipo"] = tipos[registro["tipo"]]

    server.init_db()
    conexion = server.db()
    try:
        asegurar_cxp_socios(conexion)
        cxp_socios = cuenta_id(conexion, CUENTA_CXP_SOCIOS)
        inventario_dpv = cuenta_id(conexion, "1103")
        inventario_transito = cuenta_id(conexion, "1105")

        for registro in registros:
            vehiculo = conexion.execute(
                "SELECT * FROM vehiculos WHERE vin=?", (registro["vin"],)
            ).fetchone()
            if not vehiculo:
                raise ValueError(f"No existe el vehículo maestro {registro['vin']}.")
            if conexion.execute(
                "SELECT 1 FROM adquisiciones WHERE vehiculo_id=?", (vehiculo["id"],)
            ).fetchone():
                raise ValueError(f"El VIN {registro['vin']} ya tiene una adquisición.")

            proveedor = proveedor_id(conexion, registro["proveedor"])
            nota = "Migración histórica. Contrapartida contable temporal: CxP socios."
            cur = conexion.execute(
                """INSERT INTO adquisiciones
                (vehiculo_id,proveedor_id,fecha,costo_compra,anticipo,metodo_pago,
                 condicion_pago,dias_credito,saldo,observaciones,necesita_reparacion,
                 tipo_compra,placa_cambio,banco)
                VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (vehiculo["id"], proveedor, registro["fecha"], registro["costo"], 0,
                 "Efectivo", "Contado", 0, 0, nota, 0, registro["tipo"], None, None),
            )
            conexion.execute(
                """UPDATE vehiculos
                SET precio_compra=?, proveedor=?, fecha_adquisicion=?, tipo_compra=?
                WHERE id=?""",
                (registro["costo"], registro["proveedor"], registro["fecha"],
                 registro["tipo"], vehiculo["id"]),
            )
            estado = server.recalcular_estado(
                conexion, vehiculo["id"], "Etapa calculada al registrar adquisición histórica"
            )
            server.add_movimiento(
                conexion, vehiculo["id"], registro["fecha"], "Registro de adquisición histórica",
                vehiculo["estado"], estado, vehiculo["ubicacion"], vehiculo["ubicacion"],
                referencia="CxP socios",
                observaciones=f"Compra {registro['tipo']}: {registro['proveedor']}. {nota}",
            )
            cuenta_inventario = inventario_transito if registro["tipo"] == "Importación" else inventario_dpv
            server.registrar_asiento(
                conexion, registro["fecha"],
                f"Adquisición histórica · {registro['vin']} · contrapartida CxP socios",
                "migracion_adquisicion_socios", cur.lastrowid,
                [
                    {"cuenta_id": cuenta_inventario, "debe": registro["costo"]},
                    {"cuenta_id": cxp_socios, "haber": registro["costo"]},
                ],
                vehiculo["id"],
            )
        conexion.commit()
    except Exception:
        conexion.rollback()
        raise
    finally:
        conexion.close()
        libro.close()

    print(f"Adquisiciones históricas registradas: {len(registros)}")
    print(f"Total cargado contra CxP socios: {sum(r['costo'] for r in registros):,.2f}")


if __name__ == "__main__":
    main()
