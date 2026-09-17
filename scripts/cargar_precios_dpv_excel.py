"""Carga precios de venta de la hoja Vehiculo para unidades DPV sin precio.

Se ejecuta una sola vez para completar el precio comercial sin cambiar etapas
ni generar efectos contables.
"""
from pathlib import Path
import sqlite3
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DATABASE = ROOT / "data" / "autolote.sqlite"
SOURCE = Path("/Users/erickmadrid/Library/Mobile Documents/com~apple~CloudDocs/Autolote AAA/Sistema de Gestión Corporación Triple AAA V3 Version WEb.xlsm")


def main():
    workbook = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True, keep_vba=True)
    sheet = workbook["Vehiculo"]
    headers = next(sheet.iter_rows(min_row=5, max_row=5, values_only=True))
    index = {str(value).strip(): position for position, value in enumerate(headers) if value is not None}
    prices = {
        str(row[index["VIN"]]).strip().upper(): round(float(row[index["Precio_Venta"]]), 2)
        for row in sheet.iter_rows(min_row=6, values_only=True)
        if row[index["VIN"]]
        and row[index["Precio_Venta"]] not in (None, "")
        and float(row[index["Precio_Venta"]] or 0) > 0
    }
    connection = sqlite3.connect(DATABASE)
    connection.row_factory = sqlite3.Row
    pending = connection.execute("""
        SELECT id, vin, estado, precio_venta, ubicacion
        FROM vehiculos
        WHERE estado='DPV' AND COALESCE(precio_venta, 0)<=0
        ORDER BY vin
    """).fetchall()
    missing = [row["vin"] for row in pending if row["vin"].upper() not in prices]
    if missing:
        raise RuntimeError("No se actualizó nada. VIN sin precio en Excel: " + ", ".join(missing))
    updated = []
    try:
        for row in pending:
            price = prices[row["vin"].upper()]
            connection.execute("""
                UPDATE vehiculos SET precio_venta=?
                WHERE id=? AND estado='DPV' AND COALESCE(precio_venta, 0)<=0
            """, (price, row["id"]))
            connection.execute("""
                INSERT INTO movimientos_vehiculo
                (vehiculo_id, fecha, tipo, estado_anterior, estado_nuevo,
                 ubicacion_anterior, ubicacion_nueva, referencia, observaciones)
                VALUES (?, '2026-09-15', 'Precio de venta actualizado', 'DPV', 'DPV',
                        ?, ?, 'Migración Excel', ?)
            """, (
                row["id"], row["ubicacion"], row["ubicacion"],
                f"Precio de venta cargado desde hoja Vehiculo: {price:.2f}",
            ))
            updated.append((row["vin"], price))
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()
    print(f"Actualizados: {len(updated)}")
    for vin, price in updated:
        print(f"{vin}: {price:.2f}")


if __name__ == "__main__":
    main()
