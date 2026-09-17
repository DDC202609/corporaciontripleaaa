#!/usr/bin/env python3
"""Importa exclusivamente la pestaña maestra de vehículos.

No registra adquisiciones, costos, órdenes de trabajo, ventas ni contabilidad.
"""
import os
import sys
from datetime import datetime

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
import server

ARCHIVO = "/Users/erickmadrid/Library/Mobile Documents/com~apple~CloudDocs/Autolote AAA/Vehiculos.xlsx"
HOJA = "Vehiculo"

# El VIN en tránsito conserva el dato maestro, pero su placa debe ser única.
PLACAS_CORREGIDAS = {"4A4AP3AUXEE003459": "HBL9956-1"}


def texto(valor):
    if valor is None:
        return ""
    return str(valor).strip()


def main():
    libro = load_workbook(ARCHIVO, read_only=True, data_only=True)
    hoja = libro[HOJA]
    encabezados = {texto(c.value): i for i, c in enumerate(hoja[5])}
    requeridos = ["VIN", "Lote", "Tipo_Vehiculo", "Marca", "Modelo", "Version", "Color", "año", "Placa"]
    faltantes = [nombre for nombre in requeridos if nombre not in encabezados]
    if faltantes:
        raise ValueError(f"Faltan columnas requeridas: {', '.join(faltantes)}")

    registros = []
    for fila in hoja.iter_rows(min_row=6, values_only=True):
        vin = texto(fila[encabezados["VIN"]]).upper()
        if not vin:
            continue
        placa = PLACAS_CORREGIDAS.get(vin, texto(fila[encabezados["Placa"]]).upper())
        # "N/A" no es una placa y debe conservarse como campo vacío.
        if placa in {"N/A", "NA", "-"}:
            placa = ""
        registros.append({
            "vin": vin,
            "lote": texto(fila[encabezados["Lote"]]),
            "tipo": texto(fila[encabezados["Tipo_Vehiculo"]]),
            "marca": texto(fila[encabezados["Marca"]]),
            "modelo": texto(fila[encabezados["Modelo"]]),
            "version": texto(fila[encabezados["Version"]]),
            "color": texto(fila[encabezados["Color"]]),
            "anio": fila[encabezados["año"]],
            "placa": placa,
        })

    vins = [r["vin"] for r in registros]
    placas = [r["placa"] for r in registros if r["placa"]]
    if len(vins) != len(set(vins)):
        raise ValueError("El archivo contiene VIN duplicados.")
    if len(placas) != len(set(placas)):
        raise ValueError("El archivo contiene placas duplicadas después de aplicar la corrección.")

    server.init_db()
    conexion = server.db()
    existentes = conexion.execute("SELECT COUNT(*) FROM vehiculos").fetchone()[0]
    if existentes:
        raise RuntimeError(f"La base ya tiene {existentes} vehículos; no se importó nada.")

    fecha = datetime.now().strftime("%Y-%m-%d")
    try:
        for r in registros:
            anio = int(r["anio"]) if r["anio"] not in (None, "") else None
            cur = conexion.execute(
                """INSERT INTO vehiculos
                (vin,lote,tipo_vehiculo,marca,modelo,version,color,anio,kilometraje,placa,
                 estado,ubicacion,precio_compra,precio_venta,fecha_adquisicion,proveedor,
                 tipo_compra,observaciones)
                VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (r["vin"], r["lote"], r["tipo"], r["marca"], r["modelo"], r["version"],
                 r["color"], anio, None, r["placa"], "En Tránsito", None, 0, 0,
                 None, None, None, "Dato maestro importado desde Vehiculos.xlsx"),
            )
            server.add_movimiento(
                conexion, cur.lastrowid, fecha, "Ingreso a inventario",
                estado_nuevo="En Tránsito",
                observaciones="Dato maestro importado; pendiente de adquisición",
            )
        conexion.commit()
    except Exception:
        conexion.rollback()
        raise
    finally:
        conexion.close()
        libro.close()

    print(f"Vehículos maestros importados: {len(registros)}")
    print("Placa corregida: 4A4AP3AUXEE003459 -> HBL9956-1")


if __name__ == "__main__":
    main()
